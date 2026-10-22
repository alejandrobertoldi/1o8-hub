#!/usr/bin/env python3
"""
Amazon PPC Keyword Research — Excel Workbook Generator

Takes classified keyword data and produces a professionally formatted Excel workbook
with Master Keyword List, Single Keyword Campaigns, Negative Keywords, and Strategy Summary.

Usage: This script is called by the skill after keyword classification is complete.
       It expects structured data passed via JSON stdin or imported as a module.
"""

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ── Style constants ──────────────────────────────────────────────────────────

HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
DATA_FONT = Font(name='Arial', size=10)
BOLD_FONT = Font(name='Arial', size=10, bold=True)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Color pairs: (header_fill, row_fill)
GROUP_COLORS = [
    (PatternFill('solid', fgColor='2F5496'), PatternFill('solid', fgColor='D6E4F0')),  # Blue
    (PatternFill('solid', fgColor='548235'), PatternFill('solid', fgColor='E2EFDA')),  # Green
    (PatternFill('solid', fgColor='BF8F00'), PatternFill('solid', fgColor='FFF2CC')),  # Orange
    (PatternFill('solid', fgColor='2E75B6'), PatternFill('solid', fgColor='DAEEF3')),  # Teal
    (PatternFill('solid', fgColor='7030A0'), PatternFill('solid', fgColor='E4DFEC')),  # Purple
    (PatternFill('solid', fgColor='C00000'), PatternFill('solid', fgColor='FCE4EC')),  # Red
]

NEG_HEADER_FILL = PatternFill('solid', fgColor='C00000')
NEG_CAT_FILL = PatternFill('solid', fgColor='FCE4EC')


def priority_label(sv):
    if sv >= 5000: return 'High'
    if sv >= 1000: return 'Medium'
    if sv >= 300: return 'Low'
    return 'Very Low'


def campaign_type(sv):
    if sv >= 10000: return 'Exact + Broad'
    if sv >= 3000: return 'Exact'
    return 'Exact (Niche)'


def write_kw_block(ws, start_row, keywords, group_name, group_idx):
    """Write a semantic group block. Returns next available row."""
    header_fill, row_fill = GROUP_COLORS[group_idx % len(GROUP_COLORS)]
    headers = ['Keyword', 'Search Volume', 'Cerebro IQ Score', 'PPC Sugg. Bid',
               'Competing Products', 'Title Density', 'Priority']

    # Group title row
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(headers))
    cell = ws.cell(row=start_row, column=1, value=group_name)
    cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    start_row += 1

    # Column headers
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    start_row += 1

    # Data rows
    for kw in keywords:
        sv = kw.get('search_volume', 0)
        row_data = [
            kw['keyword'],
            int(sv),
            kw.get('cerebro_iq', ''),
            kw.get('ppc_bid', ''),
            kw.get('competing_products', ''),
            kw.get('title_density', ''),
            priority_label(sv)
        ]
        for i, val in enumerate(row_data, 1):
            cell = ws.cell(row=start_row, column=i, value=val)
            cell.font = DATA_FONT
            cell.fill = row_fill
            cell.alignment = LEFT_ALIGN if i == 1 else CENTER_ALIGN
            cell.border = THIN_BORDER
        start_row += 1

    return start_row + 1  # blank row between groups


def build_workbook(config):
    """
    Build the keyword research Excel workbook.

    config dict keys:
      - product_title: str
      - product_subtitle: str
      - asin: str
      - semantic_groups: list of {name, keywords: [{keyword, search_volume, cerebro_iq, ppc_bid, competing_products, title_density}]}
      - single_keywords: list of keyword dicts (same schema as above)
      - negative_keywords: dict of {category_name: [{keyword, reason, match_type}]}
      - strategy_summary: list of (label, value) tuples
      - output_path: str
    """
    wb = Workbook()

    # ── Sheet 1: Master Keyword List ──
    ws1 = wb.active
    ws1.title = "Master Keyword List"
    ws1.column_dimensions['A'].width = 45
    ws1.column_dimensions['B'].width = 16
    ws1.column_dimensions['C'].width = 18
    ws1.column_dimensions['D'].width = 16
    ws1.column_dimensions['E'].width = 20
    ws1.column_dimensions['F'].width = 14
    ws1.column_dimensions['G'].width = 12

    ws1.merge_cells('A1:G1')
    cell = ws1.cell(row=1, column=1, value=config['product_title'])
    cell.font = Font(name='Arial', bold=True, size=14, color='2F5496')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    ws1.merge_cells('A2:G2')
    cell = ws1.cell(row=2, column=1, value=config['product_subtitle'])
    cell.font = Font(name='Arial', size=10, italic=True, color='808080')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    row = 4
    for idx, group in enumerate(config['semantic_groups']):
        row = write_kw_block(ws1, row, group['keywords'], group['name'], idx)

    ws1.freeze_panes = 'A4'

    # ── Sheet 2: Single Keyword Campaigns ──
    ws2 = wb.create_sheet("Single Keyword Campaigns")
    for col, w in [('A', 45), ('B', 16), ('C', 18), ('D', 16), ('E', 20), ('F', 14), ('G', 12)]:
        ws2.column_dimensions[col].width = w

    ws2.merge_cells('A1:G1')
    cell = ws2.cell(row=1, column=1, value="Single Keyword Campaigns — High Priority Targets")
    cell.font = Font(name='Arial', bold=True, size=14, color='2F5496')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    ws2.merge_cells('A2:G2')
    cell = ws2.cell(row=2, column=1, value="These keywords deserve their own dedicated campaign due to high volume and strong product relevance")
    cell.font = Font(name='Arial', size=10, italic=True, color='808080')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    headers = ['Keyword', 'Search Volume', 'Cerebro IQ Score', 'PPC Sugg. Bid',
               'Competing Products', 'Title Density', 'Campaign Type']
    blue_fill = PatternFill('solid', fgColor='2F5496')
    light_blue = PatternFill('solid', fgColor='D6E4F0')

    for i, h in enumerate(headers, 1):
        cell = ws2.cell(row=4, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = blue_fill
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

    r = 5
    for kw in config['single_keywords']:
        sv = kw.get('search_volume', 0)
        vals = [kw['keyword'], int(sv), kw.get('cerebro_iq', ''), kw.get('ppc_bid', ''),
                kw.get('competing_products', ''), kw.get('title_density', ''), campaign_type(sv)]
        for i, val in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=i, value=val)
            cell.font = DATA_FONT
            cell.fill = light_blue
            cell.alignment = LEFT_ALIGN if i == 1 else CENTER_ALIGN
            cell.border = THIN_BORDER
        r += 1

    ws2.freeze_panes = 'A5'

    # ── Sheet 3: Negative Keywords ──
    ws3 = wb.create_sheet("Negative Keywords")
    ws3.column_dimensions['A'].width = 25
    ws3.column_dimensions['B'].width = 55
    ws3.column_dimensions['C'].width = 20

    ws3.merge_cells('A1:C1')
    cell = ws3.cell(row=1, column=1, value=f"Negative Keywords — {config.get('asin', '')} PPC Campaigns")
    cell.font = Font(name='Arial', bold=True, size=14, color='C00000')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    ws3.merge_cells('A2:C2')
    cell = ws3.cell(row=2, column=1, value="Add these as negative phrase/exact match to prevent wasted ad spend")
    cell.font = Font(name='Arial', size=10, italic=True, color='808080')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    neg_headers = ['Negative Keyword', 'Reason', 'Match Type']
    for i, h in enumerate(neg_headers, 1):
        cell = ws3.cell(row=4, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = NEG_HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

    r = 5
    for cat_name, cat_kws in config['negative_keywords'].items():
        ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        cell = ws3.cell(row=r, column=1, value=cat_name)
        cell.font = Font(name='Arial', bold=True, size=11, color='C00000')
        cell.fill = NEG_CAT_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center')
        r += 1
        for nk in cat_kws:
            ws3.cell(row=r, column=1, value=nk['keyword']).font = DATA_FONT
            ws3.cell(row=r, column=1).border = THIN_BORDER
            ws3.cell(row=r, column=2, value=nk['reason']).font = DATA_FONT
            ws3.cell(row=r, column=2).border = THIN_BORDER
            ws3.cell(row=r, column=3, value=nk.get('match_type', 'Phrase')).font = DATA_FONT
            ws3.cell(row=r, column=3).alignment = CENTER_ALIGN
            ws3.cell(row=r, column=3).border = THIN_BORDER
            r += 1
        r += 1

    ws3.freeze_panes = 'A5'

    # ── Sheet 4: Strategy Summary ──
    ws4 = wb.create_sheet("Strategy Summary")
    ws4.column_dimensions['A'].width = 35
    ws4.column_dimensions['B'].width = 70

    ws4.merge_cells('A1:B1')
    cell = ws4.cell(row=1, column=1, value="Keyword Research Strategy Summary")
    cell.font = Font(name='Arial', bold=True, size=14, color='2F5496')
    cell.alignment = Alignment(horizontal='center', vertical='center')

    r = 3
    for label, value in config['strategy_summary']:
        ws4.cell(row=r, column=1, value=label).font = BOLD_FONT if label else DATA_FONT
        cell_b = ws4.cell(row=r, column=2, value=value)
        cell_b.font = DATA_FONT
        cell_b.alignment = Alignment(wrap_text=True)
        r += 1

    # Save
    wb.save(config['output_path'])
    return config['output_path']


if __name__ == '__main__':
    if len(sys.argv) > 1:
        with open(sys.argv[1]) as f:
            config = json.load(f)
    else:
        config = json.load(sys.stdin)

    path = build_workbook(config)
    print(f"Workbook saved to: {path}")
